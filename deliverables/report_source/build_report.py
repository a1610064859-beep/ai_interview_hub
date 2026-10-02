"""Build the editable report from its reviewed Markdown source.

Use the Codex bundled Python runtime. PDF export and page QA are separate
steps performed with the documents skill renderer.
"""
from pathlib import Path
import json
import re
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
OUT = HERE.parent
SOURCE = HERE / "项目报告书_命题05.md"
DOCX = OUT / "项目报告书_命题05.docx"
FIG = HERE / "figures"
FIG.mkdir(exist_ok=True)
FONT = "C:/Windows/Fonts/msyh.ttc"
BOLD = "C:/Windows/Fonts/msyhbd.ttc"

def fnt(size, bold=False):
    return ImageFont.truetype(BOLD if bold else FONT, size)

def centered(draw, box, lines, size=30, bold=False, color="#14202F"):
    if isinstance(lines, str):
        lines = lines.split("\n")
    line_h = size + 15
    y = (box[1] + box[3] - len(lines) * line_h) / 2
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=fnt(size, bold))
        x = (box[0] + box[2] - bbox[2] + bbox[0]) / 2
        draw.text((x, y), line, font=fnt(size, bold), fill=color)
        y += line_h

def box(draw, bounds, text, fill="#EEF3F8", size=30, bold=False):
    draw.rounded_rectangle(bounds, radius=12, fill=fill, outline="#758397", width=3)
    centered(draw, bounds, text, size, bold)

def arrow(draw, points, color="#324A66"):
    draw.line(points, fill=color, width=5, joint="curve")
    x, y = points[-1]
    px, py = points[-2]
    if abs(x - px) >= abs(y - py):
        sign = 1 if x > px else -1
        tri = [(x, y), (x - sign * 15, y - 9), (x - sign * 15, y + 9)]
    else:
        sign = 1 if y > py else -1
        tri = [(x, y), (x - 9, y - sign * 15), (x + 9, y - sign * 15)]
    draw.polygon(tri, fill=color)

def architecture():
    im = Image.new("RGB", (1600, 970), "white")
    d = ImageDraw.Draw(im)
    centered(d, (0, 10, 1600, 90), "浏览器入口与服务端驱动的训练数据链路", 39, True)
    boxes = [(45, 130, 475, 250), (585, 130, 1015, 250), (1125, 130, 1555, 250)]
    labels = ["学生端\n文本 语音 报告 成长", "新生端\n岗位路径 学习 逐题反馈", "企业端\nJD 题库 候选 原报告"]
    for b, text in zip(boxes, labels):
        box(d, b, text, size=29, bold=True)
        arrow(d, [((b[0]+b[2])//2, 250), ((b[0]+b[2])//2, 315)])
    box(d, (45, 320, 1555, 425), "Next.js 前端  录音与电平采样  图表与结果展示", size=33)
    arrow(d, [(800, 425), (800, 480)])
    box(d, (45, 485, 1555, 610), "FastAPI 服务端  题目绑定  状态编排  评分校验  历史与候选聚合", size=31, bold=True)
    for x in (260, 800, 1340):
        arrow(d, [(x, 610), (x, 670)])
    box(d, (45, 675, 475, 790), "本地语音处理\nFFmpeg  FunASR", size=30)
    box(d, (585, 675, 1015, 790), "数据保存与读取\nSQLAlchemy  SQLite", size=30)
    box(d, (1125, 675, 1555, 790), "外部或本地服务\nLLM调用链  edge-tts", size=30)
    centered(d, (40, 810, 1560, 960), ["学生与岗位 → 会话 → 回答 → 报告", "成长与企业候选复用同一报告  保留来源编号与评分口径"], 31)
    im.save(FIG / "architecture.png")

def evidence_flow():
    im = Image.new("RGB", (1600, 830), "white")
    d = ImageDraw.Draw(im)
    centered(d, (0, 0, 1600, 95), "逐题状态与最终评分各有明确边界", 38, True)
    xs = [35, 430, 825, 1220]
    labels = ["绑定当前题\n接收并保存回答", "服务端判断\n追问最多一次", "返回下一步\n追问 下一题 完成", "完成时生成报告\n失败回滚 可重试"]
    for x, label in zip(xs, labels):
        box(d, (x, 135, x+345, 270), label, size=28)
        if x != xs[-1]:
            arrow(d, [(x+345, 202), (x+390, 202)])
    centered(d, (35, 295, 1565, 365), "最终报告的语义与声学结果分别计算", 31, True)
    box(d, (35, 395, 620, 525), "语义三维  两次独立调用\n原话校验  有效值汇总", size=31)
    box(d, (980, 395, 1565, 525), "表达流畅度  声学规则\n完整样本计算  不足则未评估", size=31)
    arrow(d, [(620, 460), (680, 460), (680, 590)])
    arrow(d, [(980, 460), (920, 460), (920, 590)])
    box(d, (405, 595, 1195, 690), "保存四维报告  证据 原因 建议  有效维度均值", size=30, bold=True)
    centered(d, (20, 710, 1580, 820), "学生复盘  →  同岗可比成长  →  企业同源候选与人工复核", 30)
    im.save(FIG / "evidence_flow.png")

def font_style(style, east="宋体", size=11, bold=False):
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = False
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for key in ("ascii", "hAnsi", "eastAsia"):
        fonts.set(qn("w:"+key), east if key == "eastAsia" else "Calibri")
    for key in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme"):
        fonts.attrib.pop(qn("w:"+key), None)

def field(paragraph, code):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = code
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for item in (begin, instr, separate, text, end):
        run._r.append(item)

def bookmark(paragraph, number):
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(number))
    start.set(qn("w:name"), "chapter_" + str(number))
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(number))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)

def remove_title_borders(doc):
    for style in doc.styles:
        for border in list(style.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    for p in doc.paragraphs:
        if p.style.name == "Title":
            ppr = p._p.get_or_add_pPr()
            for border in list(ppr.findall(qn("w:pBdr"))):
                ppr.remove(border)

def table_style(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for key in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement("w:"+key)
        for attr, val in (("val", "single"), ("sz", "5"), ("color", "D9D9D9")):
            b.set(qn("w:"+attr), val)
        borders.append(b)
    pr.append(borders)
    margins = OxmlElement("w:tblCellMar")
    for key, value in (("top", "95"), ("bottom", "95"), ("left", "105"), ("right", "105")):
        m = OxmlElement("w:"+key)
        m.set(qn("w:w"), value)
        m.set(qn("w:type"), "dxa")
        margins.append(m)
    pr.append(margins)
    for i, column in enumerate(table.columns):
        column.width = Inches(widths[i])
    for ri, row in enumerate(table.rows):
        trpr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        trpr.append(cant)
        if ri == 0:
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            trpr.append(repeat)
        for ci, cell in enumerate(row.cells):
            cell.width = Inches(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "DAE5EF" if ri == 0 else ("F5F7FA" if ri % 2 == 0 else "FFFFFF"))
            cell._tc.get_or_add_tcPr().append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.first_line_indent = Pt(0)
                p.paragraph_format.line_spacing = 1.2
                p.paragraph_format.widow_control = True
                p.paragraph_format.keep_with_next = ri in (0, 1, len(table.rows) - 2)
                for run in p.runs:
                    run.font.size = Pt(10)
                    run.font.bold = ri == 0

def add_table(doc, rows):
    cols = len(rows[0])
    if cols == 3:
        widths = [1.35, 2.6, 2.55]
    elif cols == 4:
        widths = [1.14, 1.73, 1.67, 1.96]
    elif cols == 5:
        widths = [1.0, 1.1, 1.3, 1.6, 1.5]
    elif cols == 6:
        widths = [0.8, 1.0, 1.2, 1.1, 1.2, 1.2]
    elif cols == 8:
        widths = [0.8, 0.9, 0.7, 0.8, 0.8, 0.9, 0.8, 0.8]
    elif cols == 9:
        widths = [0.65, 0.8, 0.75, 0.65, 0.75, 0.75, 0.85, 0.65, 0.65]
    else:
        widths = [6.5 / cols] * cols
    if rows[0][0] == "核心接口":
        widths = [2.85, 1.43, 2.22]
    if rows[0][0] == "人员":
        widths = [0.95, 1.75, 3.8]
    table = doc.add_table(rows=len(rows), cols=cols)
    for ri, row in enumerate(rows):
        for ci, text in enumerate(row):
            table.cell(ri, ci).text = text
    table_style(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def main():
    architecture()
    evidence_flow()
    content = SOURCE.read_text(encoding="utf-8")
    lines = content.splitlines()
    titles = [ln[2:] for ln in lines if re.match(r"# (第[一二三四五六七八九十]+章|附录[AB])", ln)]
    tocfile = HERE / "toc_pages.json"
    toc = json.loads(tocfile.read_text(encoding="utf-8")) if tocfile.exists() else {}
    doc = Document()
    doc.core_properties.title = "智能面试仓的开发与应用 完整项目报告书"
    doc.core_properties.subject = "企业命题05 本科组 芝麻队"
    doc.core_properties.author = "芝麻队"
    doc.core_properties.keywords = "智能汽车 可解释评分 成长追踪"
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(0.85)
    sec.left_margin = sec.right_margin = Inches(1.0)
    sec.header_distance = sec.footer_distance = Inches(0.35)
    sec.different_first_page_header_footer = True
    normal = doc.styles["Normal"]
    font_style(normal)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.first_line_indent = Pt(22)
    normal.paragraph_format.widow_control = True
    for name, size in (("Title", 22), ("Subtitle", 15), ("Heading 1", 17), ("Heading 2", 12.5)):
        font_style(doc.styles[name], "黑体", size, name != "Subtitle")
        doc.styles[name].paragraph_format.first_line_indent = Pt(0)
        doc.styles[name].paragraph_format.keep_with_next = True
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(0)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(15)
    doc.styles["Heading 2"].paragraph_format.space_before = Pt(12)
    doc.styles["Heading 2"].paragraph_format.space_after = Pt(7)
    font_style(doc.styles["Caption"], "宋体", 9.5)
    doc.styles["Caption"].paragraph_format.first_line_indent = Pt(0)
    doc.styles["Caption"].paragraph_format.line_spacing = 1.2
    doc.styles["Caption"].paragraph_format.space_after = Pt(10)
    hp = sec.header.paragraphs[0]
    hp.paragraph_format.first_line_indent = Pt(0)
    hp.add_run("智能面试仓的开发与应用  芝麻队  项目报告书").font.size = Pt(8)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.first_line_indent = Pt(0)
    fp.add_run("第 ").font.size = Pt(9)
    field(fp, "PAGE")
    fp.add_run(" 页").font.size = Pt(9)
    in_cover, chapter, i = True, 0, 0
    table_count = {}
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line == "<!-- PAGEBREAK -->":
            doc.add_page_break()
            in_cover = False
        elif line == "<!-- TOC -->":
            toc_heading = doc.add_paragraph("目录", "Heading 1")
            toc_heading.paragraph_format.page_break_before = True
            for index, title in enumerate(titles, 1):
                p = doc.add_paragraph()
                p.paragraph_format.first_line_indent = Pt(0)
                p.paragraph_format.space_after = Pt(15)
                p.paragraph_format.tab_stops.add_tab_stop(Inches(6.45), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
                p.add_run(title + "\t" + str(toc.get(title, "—")))
            p = doc.add_paragraph("十章内容对应命题05要求，页码以本版PDF核对。")
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_before = Pt(12)
        elif line.startswith("# "):
            title = line[2:]
            if in_cover:
                p = doc.add_paragraph(title, "Title")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(92)
                p.paragraph_format.space_after = Pt(22)
            else:
                chapter += 1
                p = doc.add_paragraph(title, "Heading 1")
                p.paragraph_format.page_break_before = True
                bookmark(p, chapter)
        elif line.startswith("## "):
            doc.add_paragraph(line[3:], "Heading 2" if chapter else "Heading 1")
        elif line.startswith("!["):
            match = re.match(r"!\[(.*?)\]\((.*?)\)", line)
            caption, relative = match.groups()
            path = (HERE / relative).resolve()
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.keep_with_next = True
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(path), width=Inches(6.5))
            p = doc.add_paragraph(caption, "Caption")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [v.strip() for v in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", v) for v in cells):
                    rows.append(cells)
                i += 1
            table_count[chapter] = table_count.get(chapter, 0) + 1
            if chapter == 11:
                cap = f"附录表 A-{table_count[chapter]}  {rows[0][0]}对应表"
            elif chapter == 12:
                cap = f"附录表 B-{table_count[chapter]}  {rows[0][0]}对应表"
            else:
                cap = "表 " + str(chapter) + " " + str(table_count[chapter]) + "  " + rows[0][0] + "对应表"
            p = doc.add_paragraph(cap, "Caption")
            p.paragraph_format.keep_with_next = True
            add_table(doc, rows)
            continue
        else:
            if in_cover:
                p = doc.add_paragraph(line)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.first_line_indent = Pt(0)
                p.paragraph_format.line_spacing = 1.4
                p.paragraph_format.space_after = Pt(9)
                if line.startswith("版本说明"):
                    p.paragraph_format.space_before = Pt(18)
                    for run in p.runs:
                        run.font.size = Pt(9)
                if line == "智驾未来 AI面试仓":
                    p.style = doc.styles["Subtitle"]
            elif line.startswith("［") or line.startswith("另："):
                p = doc.add_paragraph(line)
                p.paragraph_format.first_line_indent = Pt(0)
                p.paragraph_format.line_spacing = 1.2
                p.paragraph_format.space_after = Pt(4)
                for run in p.runs:
                    run.font.size = Pt(10.5)
            else:
                p = doc.add_paragraph()
                if "**" in line:
                    parts = re.split(r"(\*\*.*?\*\*)", line)
                    for part in parts:
                        if part.startswith("**") and part.endswith("**") and len(part) >= 4:
                            r = p.add_run(part[2:-2])
                            r.bold = True
                        else:
                            p.add_run(part)
                else:
                    p.add_run(line)
        i += 1
    remove_title_borders(doc)
    doc.save(DOCX)
    metrics = {
        "title": doc.core_properties.title,
        "chapter_count": len(titles),
        "character_count": len(content),
        "chinese_character_count": len(re.findall(r"[\u4e00-\u9fff]", content)),
        "table_count": len(doc.tables),
        "figure_count": len(doc.inline_shapes),
        "docx": str(DOCX),
        "toc_pages": toc,
    }
    (HERE / "build_metrics.json").write_text(json.dumps(metrics, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(metrics, ensure_ascii=False))

if __name__ == "__main__":
    main()
